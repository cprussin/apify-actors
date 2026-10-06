"""Writes tiny test documents to the directory given as argv[1].

Generated at test time (nothing binary is committed). Needs the actor's
Python requirements (python-pptx, openpyxl and Pillow come with markitdown).
"""

import os
import sys
import zipfile

out = sys.argv[1]
os.makedirs(out, exist_ok=True)


def write_zip(name, entries, first=None):
    with zipfile.ZipFile(os.path.join(out, name), "w", zipfile.ZIP_DEFLATED) as z:
        if first:
            z.writestr(zipfile.ZipInfo(first[0]), first[1], zipfile.ZIP_STORED)
        for path, data in entries.items():
            z.writestr(path, data)


def pdf(name, pages):
    """A minimal text PDF: one Helvetica text block per page."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for lines in pages:
        ops = "BT /F1 14 Tf 72 720 Td 18 TL " + " ".join(
            f"({line}) Tj T*" for line in lines
        ) + " ET"
        objs.append(f"<< /Length {len(ops)} >>\nstream\n{ops}\nendstream")
        content = len(objs)
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content} 0 R >>"
        )
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    body = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(body))
        body += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(body)
    body += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    body += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    body += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    with open(os.path.join(out, name), "wb") as f:
        f.write(body)


pdf(
    "text.pdf",
    [
        ["Quarterly report", "Revenue grew in every region this quarter."],
        ["Second page", "Costs were flat compared with last year."],
    ],
)

# A scanned page: an image of text with no text layer.
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

img = Image.new("L", (1700, 600), 255)
draw = ImageDraw.Draw(img)
font = ImageFont.load_default(size=72)
draw.text((80, 120), "Scanned invoice", fill=0, font=font)
draw.text((80, 300), "Total due 1250 dollars", fill=0, font=font)
img.save(os.path.join(out, "scanned.pdf"), resolution=200)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def p(text, style=None):
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{ppr}<w:r><w:t>{text}</w:t></w:r></w:p>"


def cell(text):
    return f"<w:tc>{p(text)}</w:tc>"


table = "<w:tbl>" + "".join(
    "<w:tr>" + "".join(cell(c) for c in row) + "</w:tr>"
    for row in [["Region", "Revenue"], ["EU", "120"], ["US", "340"]]
) + "</w:tbl>"
write_zip(
    "table.docx",
    {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>"
        ),
        "word/document.xml": (
            f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{W}"><w:body>'
            + p("Sales summary", "Heading1")
            + p("Revenue by region:")
            + table
            + "</w:body></w:document>"
        ),
    },
)

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

prs = Presentation()
for i in range(2):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = f"Slide title {i + 1}"
    s.placeholders[1].text = f"Point {i + 1}"
t = prs.slides[1].shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
for (r, c), v in {(0, 0): "Plan", (0, 1): "Price", (1, 0): "Pro", (1, 1): "9"}.items():
    t.cell(r, c).text = v
prs.save(os.path.join(out, "deck.pptx"))

import openpyxl  # noqa: E402

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sales"
ws.append(["Region", "Total"])
ws.append(["EU", 5])
ws2 = wb.create_sheet("Costs")
ws2.append(["Item", "Amount"])
ws2.append(["Rent", 2])
wb.save(os.path.join(out, "book.xlsx"))

with open(os.path.join(out, "page.html"), "w", encoding="utf-8") as f:
    f.write(
        "<!doctype html><html><head><title>Price list</title></head><body>"
        "<h1>Prices</h1><p>All prices in <b>EUR</b>.</p>"
        "<table><tr><th>Item</th><th>Price</th></tr><tr><td>Tea</td><td>3</td></tr></table>"
        "</body></html>"
    )

write_zip(
    "book.epub",
    {
        "META-INF/container.xml": (
            '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>'
        ),
        "content.opf": (
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Tiny Book</dc:title>'
            '<dc:creator>Ann Author</dc:creator><dc:identifier id="id">x</dc:identifier></metadata>'
            '<manifest><item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/></manifest>'
            '<spine><itemref idref="c1"/></spine></package>'
        ),
        "ch1.xhtml": (
            '<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter 1</title></head>'
            "<body><h1>Chapter One</h1><p>It was a quiet morning.</p></body></html>"
        ),
    },
    first=("mimetype", "application/epub+zip"),
)
